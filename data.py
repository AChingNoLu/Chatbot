import json
import re
import subprocess
import tempfile
from pathlib import Path

from docx import Document


# =========================
# CẤU HÌNH
# =========================

BASE_DIR = Path(__file__).parent
DATA_DIR = BASE_DIR / "DATA"
OUTPUT_FILE = BASE_DIR / "dataset_rag_final.json"


# =========================
# XÁC ĐỊNH LOẠI VĂN BẢN
# =========================

def get_loai_van_ban(file_path):

    folder = file_path.parent.name.lower()

    if folder == "luat":
        return "Luật"

    if folder == "nghi_dinh":
        return "Nghị định"

    if folder == "thong_tu":
        return "Thông tư"

    if folder == "danhmucnganhnghe":
        return "Danh mục ngành nghề"

    return "Khác"


# =========================
# ĐỌC FILE DOC
# =========================

def convert_doc_to_docx(doc_file):

    temp_dir = Path(tempfile.mkdtemp())

    try:

        subprocess.run(
            [
                "soffice",
                "--headless",
                "--convert-to",
                "docx",
                "--outdir",
                str(temp_dir),
                str(doc_file)
            ],
            check=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE
        )

        new_file = temp_dir / (doc_file.stem + ".docx")

        if new_file.exists():
            return new_file

    except Exception as e:

        print(f"Không thể chuyển {doc_file.name}: {e}")

    return None


# =========================
# ĐỌC DOCX
# =========================

def read_docx(file_path):

    doc = Document(file_path)

    lines = []

    # Đọc đoạn văn
    for paragraph in doc.paragraphs:

        text = paragraph.text.strip()

        if text:
            lines.append(text)

    # Đọc bảng
    for table in doc.tables:

        for row in table.rows:

            cells = []

            for cell in row.cells:

                text = cell.text.strip()
                cells.append(text)

            if any(cells):

                lines.append(
                    " | ".join(cells)
                )

    return "\n".join(lines)


# =========================
# ĐỌC DOC / DOCX
# =========================

def read_file(file_path):

    suffix = file_path.suffix.lower()

    if suffix == ".docx":
        return read_docx(file_path)

    if suffix == ".doc":

        docx_file = convert_doc_to_docx(file_path)

        if docx_file:
            return read_docx(docx_file)

        return ""

    return ""


# =========================
# LẤY SỐ HIỆU
# =========================

def get_so_hieu(text, file_name):

    patterns = [

        r"\b\d+/\d{4}/NĐ-CP\b",

        r"\b\d+/\d{4}/TT-[A-ZĐ\-]+\b",

        r"\b\d+/VBHN-VPQH\b",

        r"\b\d+/\d{4}/QH\d+\b"

    ]

    for pattern in patterns:

        match = re.search(
            pattern,
            text
        )

        if match:
            return match.group(0)

    for pattern in patterns:

        match = re.search(
            pattern,
            file_name
        )

        if match:
            return match.group(0)

    return ""


# =========================
# LẤY TÊN VĂN BẢN
# =========================

def get_ten_van_ban(text, file_name):

    lines = text.split("\n")

    for line in lines[:30]:

        line = line.strip()

        if len(line) > 20:

            # Bỏ các dòng quá giống tiêu đề cơ quan
            if "CỘNG HÒA" not in line.upper():

                return line

    return Path(file_name).stem


# =========================
# LẤY NGÀY CÓ HIỆU LỰC
# =========================

def get_ngay_co_hieu_luc(text):

    # Ví dụ:
    # Có hiệu lực từ ngày 01/02/2021
    # Có hiệu lực kể từ ngày 01 tháng 02 năm 2021
    # Có hiệu lực thi hành từ ngày 01 tháng 01 năm 2021

    patterns = [

        r"[Cc]ó hiệu lực.*?ngày\s+(\d{1,2})[\/\-](\d{1,2})[\/\-](\d{4})",

        r"[Cc]ó hiệu lực.*?ngày\s+(\d{1,2})\s+tháng\s+(\d{1,2})\s+năm\s+(\d{4})",

        r"[Cc]ó hiệu lực thi hành.*?ngày\s+(\d{1,2})\s+tháng\s+(\d{1,2})\s+năm\s+(\d{4})",

        r"[Cc]ó hiệu lực kể từ ngày\s+(\d{1,2})[\/\-](\d{1,2})[\/\-](\d{4})"
    ]

    for pattern in patterns:

        match = re.search(
            pattern,
            text,
            re.IGNORECASE | re.DOTALL
        )

        if match:

            day = match.group(1).zfill(2)
            month = match.group(2).zfill(2)
            year = match.group(3)

            return f"{day}/{month}/{year}"

    return ""


# =========================
# CHIA THEO ĐIỀU
# =========================

def split_by_dieu(text):

    pattern = r"(?=Điều\s+\d+\s*[\.:])"

    parts = re.split(
        pattern,
        text
    )

    chunks = []

    for part in parts:

        part = part.strip()

        if len(part) < 30:
            continue

        match = re.match(
            r"(Điều\s+\d+)",
            part
        )

        if match:
            dieu = match.group(1)
        else:
            dieu = ""

        chunks.append({
            "dieu": dieu,
            "content": part
        })

    return chunks


# =========================
# TẠO DATASET
# =========================

def build_dataset():

    dataset = []

    # Tìm tất cả DOC và DOCX
    files = list(DATA_DIR.rglob("*.doc"))
    files += list(DATA_DIR.rglob("*.docx"))

    print(f"Tìm thấy {len(files)} file văn bản.")

    for file_path in files:

        print(
            f"\nĐang xử lý: {file_path.name}"
        )

        text = read_file(file_path)

        if not text.strip():

            print("  -> Không đọc được nội dung")
            continue

        loai_van_ban = get_loai_van_ban(
            file_path
        )

        so_hieu = get_so_hieu(
            text,
            file_path.name
        )

        ten_van_ban = get_ten_van_ban(
            text,
            file_path.name
        )

        ngay_co_hieu_luc = get_ngay_co_hieu_luc(
            text
        )

        # Nếu không tìm được trong nội dung
        if not ngay_co_hieu_luc:

            print(
                "  -> Không tìm thấy ngày có hiệu lực"
            )

        chunks = split_by_dieu(text)

        # Nếu không chia được theo Điều
        if not chunks:

            chunks = [
                {
                    "dieu": "",
                    "content": text
                }
            ]

        # =========================
        # TẠO CHUNK
        # =========================

        for chunk in chunks:

            item = {

                "file_name": file_path.name,

                "loai_van_ban": loai_van_ban,

                "so_hieu": so_hieu,

                "ten_van_ban": ten_van_ban,

                "ngay_co_hieu_luc": ngay_co_hieu_luc,

                "hieu_luc": "Còn hiệu lực",

                "dieu": chunk["dieu"],

                "content": chunk["content"]
            }

            dataset.append(item)

    # =========================
    # LƯU JSON
    # =========================

    with open(
        OUTPUT_FILE,
        "w",
        encoding="utf-8"
    ) as f:

        json.dump(
            dataset,
            f,
            ensure_ascii=False,
            indent=2
        )

    print("\n==============================")
    print("ĐÃ TẠO DATASET")
    print("==============================")

    print(
        f"Số chunk: {len(dataset)}"
    )

    print(
        f"File: {OUTPUT_FILE}"
    )


# =========================
# MAIN
# =========================

if __name__ == "__main__":

    build_dataset()